"""Excel export of delegates and renewal plans:
 * «السيارات» sheet: «ملف المفوّض» + «انتهاء رخصة المفوّض» next to «المفوّض بالقيادة».
 * new sheet «المفوّضون»: file code, name, licence date/status, vehicles; licence no. and mobile only with view_sensitive.
 * new sheet «خطط التجديد»: plan, vehicle, delegate, progress, next step, the five steps' status, plan note.
 * nothing in the database changes. Fake data only.
Usage: python3 -I ui_excel_renewal.py OUTDIR   (fresh fleet_ui + mock server)"""
import sys, os, subprocess, datetime
import openpyxl
from playwright.sync_api import sync_playwright
OUT = sys.argv[1]; os.makedirs(OUT, exist_ok=True)
BASE = 'http://localhost:8787/'
problems, passed = [], []
def ok(c, m): (passed if c else problems).append(m)
def sql(q): return subprocess.run(['psql', '-h', '/tmp', '-p', '5433', '-U', 'postgres', '-X', '-q', '-At', '-d', 'fleet_ui', '-c', q], capture_output=True, text=True).stdout.strip()
def new_ctx(b, w=1280):
    c = b.new_context(viewport={'width': w, 'height': 900}, locale='ar-SA', timezone_id='Asia/Riyadh', accept_downloads=True)
    g = c.new_page(); g.route('**/fonts.googleapis.com/**', lambda r: r.abort())
    g.on('pageerror', lambda e: problems.append(f'page error: {e}'))
    return c, g
def login(g, email):
    g.goto(BASE); g.wait_for_selector('form.login-card'); g.fill('input[name=email]', email); g.fill('input[name=password]', 'test1234'); g.click('button[type=submit]'); g.wait_for_selector('.panels')
def download(g, sel, name):
    with g.expect_download() as dl: g.click(sel)
    path = f'{OUT}/{name}'; dl.value.save_as(path); return path
def table(ws):
    h = [c.value for c in ws[1]]
    return h, [dict(zip(h, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
V1 = 'b0000000-0000-0000-0000-000000000001'
V2 = 'b0000000-0000-0000-0000-000000000002'
# fake delegate linked to two vehicles, with licence + mobile in the private table; a plan for vehicle 1
sql(f"""set request.jwt.claim.sub = 'a0000000-0000-0000-0000-000000000001';
insert into delegates(full_name, license_expiry) values ('مفوض اختبار التصدير', app_today() + 400);
insert into delegate_private(delegate_id, license_no, mobile) select id, '1999999999', '+966500000009' from delegates where full_name = 'مفوض اختبار التصدير';
update vehicles set delegate_id = (select id from delegates where full_name = 'مفوض اختبار التصدير') where id in ('{V1}', '{V2}');
update documents set expiry_date = app_today() - 5 where vehicle_id = '{V1}' and doc_type in ('periodic_inspection', 'registration');
update vehicles set insurance_expiry = app_today() + 200 where id = '{V1}';""")
sql(f"set request.jwt.claim.sub = 'a0000000-0000-0000-0000-000000000001'; set role authenticated; select create_renewal_plan('{V1}', null);")
PID = sql(f"select id from renewal_plans where vehicle_id = '{V1}'")
sql(f"set request.jwt.claim.sub = 'a0000000-0000-0000-0000-000000000001'; set role authenticated; select save_renewal_step('{PID}', 4, '{{\"fines_status\":\"paid\"}}', false); select set_renewal_plan_note('{PID}', 'ملاحظة اختبار للخطة');")
LIC = sql("select to_char(license_expiry, 'YYYY-MM-DD') from delegates where full_name = 'مفوض اختبار التصدير'")
CODE = sql("select code from delegates where full_name = 'مفوض اختبار التصدير'")
TABLES = ['vehicles', 'delegates', 'delegate_private', 'renewal_plans', 'renewal_steps', 'documents']
FP = lambda: {t: sql(f"select md5(coalesce(string_agg(x::text, '|' order by x::text), '')) from {t} x") for t in TABLES}
before = FP()
d = lambda v: v.date().isoformat() if isinstance(v, datetime.datetime) else (v.isoformat() if isinstance(v, datetime.date) else v)

with sync_playwright() as p:
    b = p.chromium.launch()
    c, g = new_ctx(b); login(g, 'admin@test.local')
    g.goto(BASE + '#/excel'); g.wait_for_selector('[data-exp=renewal]')
    ok(g.locator('[data-exp=delegates]').count() == 1 and g.locator('[data-exp=renewal]').count() == 1, 'excel page has «المفوّضون» and «خطط التجديد» buttons')
    g.screenshot(path=f'{OUT}/excel-page.png', full_page=True)

    # vehicles sheet
    ws = openpyxl.load_workbook(download(g, '[data-exp=vehicles]', 'vehicles.xlsx')).active
    h, rows = table(ws)
    i = h.index('المفوّض بالقيادة') if 'المفوّض بالقيادة' in h else -1
    ok(i >= 0 and h[i + 1] == 'ملف المفوّض' and h[i + 2] == 'انتهاء رخصة المفوّض', f'vehicles: delegate file + licence columns after the delegate name ({h[i:i + 3]})')
    r1 = next(r for r in rows if r['الرقم الداخلي'] == 'VEH-0001')
    ok(r1['المفوّض بالقيادة'] == 'مفوض اختبار التصدير' and r1['ملف المفوّض'] == CODE and d(r1['انتهاء رخصة المفوّض']) == LIC, f'vehicles: VEH-0001 delegate, file and licence date ({r1["ملف المفوّض"]}, {r1["انتهاء رخصة المفوّض"]})')
    r3 = next(r for r in rows if r['الرقم الداخلي'] == 'VEH-0003')
    ok(not r3['ملف المفوّض'] and not r3['انتهاء رخصة المفوّض'], 'vehicles: vehicle without a delegate has empty delegate columns')
    ok(ws.cell(2, h.index('انتهاء رخصة المفوّض') + 1).number_format.lower().startswith('yyyy/mm/dd'), 'vehicles: licence date formatted YYYY/MM/DD')

    # delegates sheet (admin sees licence no. and mobile)
    ws = openpyxl.load_workbook(download(g, '[data-exp=delegates]', 'delegates.xlsx')).active
    h, rows = table(ws)
    ok(ws.title == 'المفوّضون' and ws.sheet_view.rightToLeft, 'delegates sheet named المفوّضون, right-to-left')
    dr = next((r for r in rows if r['اسم المفوّض'] == 'مفوض اختبار التصدير'), {})
    ok(dr.get('رقم الملف') == CODE and d(dr.get('انتهاء رخصة القيادة')) == LIC and dr.get('حالة الرخصة') == 'سارية' and dr.get('عدد السيارات') == 2,
       f'delegates: code, licence date, status, 2 vehicles ({dr})')
    ok('VEH-0001' in (dr.get('السيارات') or '') and 'VEH-0002' in (dr.get('السيارات') or ''), 'delegates: both vehicles listed by internal code')
    ok(dr.get('رقم الرخصة / الهوية') == '1999999999' and dr.get('الجوال') == '+966500000009', 'delegates: licence no. and mobile for view_sensitive (admin)')

    # renewal plans sheet
    ws = openpyxl.load_workbook(download(g, '[data-exp=renewal]', 'renewal.xlsx')).active
    h, rows = table(ws)
    ok(ws.title == 'خطط التجديد' and len(rows) == 1, f'renewal sheet named خطط التجديد with the one plan ({len(rows)})')
    pr = rows[0] if rows else {}
    ok(pr.get('الرقم الداخلي') == 'VEH-0001' and pr.get('المفوّض') == 'مفوض اختبار التصدير' and pr.get('حالة الخطة') == 'جارية', 'renewal: vehicle, delegate, status')
    ok(str(pr.get('رخصة المفوّض', '')).startswith('مكتملة (تلقائي)') and str(pr.get('التأمين', '')).startswith('مكتملة (تلقائي)'), f'renewal: valid licence/insurance complete automatically ({pr.get("رخصة المفوّض")}, {pr.get("التأمين")})')
    ok(pr.get('الفحص الدوري') == 'لم تبدأ' and pr.get('المخالفات والرسوم') == 'بانتظار التأكيد' and pr.get('الاستمارة') == 'لم تبدأ', 'renewal: pending / waiting steps')
    ok(pr.get('المكتمل من 5') == 2 and pr.get('الخطوة التالية') == 'الفحص الدوري', f'renewal: 2 of 5, next step inspection ({pr.get("المكتمل من 5")}, {pr.get("الخطوة التالية")})')
    ok(pr.get('ملاحظة الخطة') == 'ملاحظة اختبار للخطة', 'renewal: plan note exported')

    # «all» file has both new sheets
    wa = openpyxl.load_workbook(download(g, '[data-exp=all]', 'all.xlsx'))
    ok('المفوّضون' in wa.sheetnames and 'خطط التجديد' in wa.sheetnames, f'«all» export includes the new sheets ({wa.sheetnames})')
    # fix: the delegate page shows the licence no. and mobile to view_sensitive users (read used a wrong sort key)
    did = sql("select id from delegates where full_name = 'مفوض اختبار التصدير'")
    g.goto(BASE + f'#/delegates/{did}'); g.wait_for_timeout(1500)
    txt = g.inner_text('#page')
    open(f'{OUT}/dlg.txt', 'w').write(txt); g.screenshot(path=f'{OUT}/delegate.png', full_page=True)
    ok('1999999999' in txt and '050 000 0009' in txt, 'delegate page shows licence no. and mobile to admin')
    c.close()

    # view-only user: no licence numbers or mobiles
    c, g = new_ctx(b); login(g, 'viewer@test.local')
    g.goto(BASE + '#/excel'); g.wait_for_timeout(1500)
    if g.locator('[data-exp=delegates]').count():
        ws = openpyxl.load_workbook(download(g, '[data-exp=delegates]', 'delegates-viewer.xlsx')).active
        h, rows = table(ws)
        ok('رقم الرخصة / الهوية' not in h and 'الجوال' not in h and len(rows) >= 1, f'viewer: delegates sheet without licence no. / mobile ({h})')
    else:
        ok(g.locator('[data-exp]').count() == 0, 'viewer: no export without the export permission')
    c.close()

    # mobile layout of the excel page
    c, g = new_ctx(b, 390); login(g, 'admin@test.local')
    g.goto(BASE + '#/excel'); g.wait_for_selector('[data-exp=renewal]')
    ok(g.evaluate('document.documentElement.scrollWidth') <= 391, 'mobile: excel page has no horizontal overflow')
    g.screenshot(path=f'{OUT}/excel-mobile.png', full_page=True)
    c.close(); b.close()

ok(FP() == before, 'export changed nothing in the database')
for m in passed: print('PASS:', m)
for m in problems: print('FAIL:', m)
print(f'PASS {len(passed)} FAIL {len(problems)}')
sys.exit(1 if problems else 0)
