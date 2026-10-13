"""Phase 8: the app with a large dataset (>1000 rows per table) behind a Supabase-like 1000-row cap.
Checks that lists, totals, reports, Excel and backup see ALL rows, and measures page times.
Usage: python3 -I bulk_ui.py OUTDIR  (fleet_ui loaded with ui_seed.sql + bulk_seed.sql; mock started with default --max-rows 1000)"""
import sys, os, time, json, subprocess, datetime
from decimal import Decimal as D
from playwright.sync_api import sync_playwright
import openpyxl

OUT = sys.argv[1]; os.makedirs(OUT, exist_ok=True)
BASE = 'http://localhost:8787/'
problems, passed, timings = [], [], []
def ok(c, m): (passed if c else problems).append(m)
PSQL = ['psql', '-h', '/tmp', '-p', '5433', '-U', 'postgres', '-X', '-q', '-At', '-d', 'fleet_ui']
def sql(q): return subprocess.run(PSQL + ['-c', q], capture_output=True, text=True).stdout.strip()
def money(v): return f'{D(v):,.2f}'
def timed(label, fn):
    t0 = time.time(); fn(); dt = time.time() - t0; timings.append((label, dt)); return dt

with sync_playwright() as p:
    b = p.chromium.launch()
    c = b.new_context(viewport={'width': 390, 'height': 844}, device_scale_factor=2, locale='ar-SA', timezone_id='Asia/Riyadh', accept_downloads=True)
    pg = c.new_page(); pg.route('**/fonts.googleapis.com/**', lambda r: r.abort())
    pg.on('pageerror', lambda e: problems.append(f'page error: {e}'))
    # 4x CPU slowdown ≈ a mid-range phone
    cdp = c.new_cdp_session(pg); cdp.send('Emulation.setCPUThrottlingRate', {'rate': 4})
    pg.goto(BASE); pg.wait_for_selector('form.login-card')
    pg.fill('input[name=email]', 'admin@test.local'); pg.fill('input[name=password]', 'test1234')
    timed('login + dashboard', lambda: (pg.click('button[type=submit]'), pg.wait_for_selector('.panels')))

    nveh = sql('select count(*) from vehicles where deleted_at is null')
    timed('vehicles list (1,105)', lambda: (pg.goto(BASE + '#/vehicles'), pg.wait_for_selector('#vt')))
    ok(pg.inner_text('.page-head .count') == f'{int(nveh):,}', f'vehicles page counts all {nveh} vehicles (more than the 1000-row cap)')
    ok(pg.locator('#vt tbody tr').count() == 100 and 'يُعرض 100 من 1,105' in pg.inner_text('#vt .more-rows'), 'long list shows first 100 rows + «عرض المزيد»')
    pg.click('#vt [data-more="200"]'); pg.wait_for_timeout(300)
    ok(pg.locator('#vt tbody tr').count() == 300, '«عرض 200 أخرى» → 300 rows')
    timed('show all 1,105 vehicles', lambda: (pg.click('#vt [data-more=all]'), pg.wait_for_function("document.querySelectorAll('#vt tbody tr').length === 1105")))
    ok(pg.locator('#vt .more-rows').count() == 0, '«عرض الكل» shows every vehicle')
    pg.fill('#page input[name=q]', 'ت ج ر 10777'); pg.wait_for_timeout(300)
    ok(pg.locator('#vt tbody tr').count() == 1, 'filter searches all 1,105 vehicles (not only the visible 100)')

    timed('vehicle file', lambda: (pg.goto(BASE + '#/vehicles/b0000000-0000-0000-0000-000000000001'), pg.wait_for_selector('.vhead')))
    timed('documents & alerts', lambda: (pg.goto(BASE + '#/documents'), pg.wait_for_selector('.page-head')))
    pg.wait_for_timeout(300)

    def list_total(page, sqltable, datecol, label):
        timed(label, lambda: (pg.goto(BASE + f'#/{page}'), pg.wait_for_selector('.list-body')))
        pg.select_option('#page select[name=year]', ''); pg.wait_for_timeout(400)
        want = sql(f"select count(*)||'|'||sum(total) from {sqltable} where deleted_at is null")
        n, s = want.split('|')
        t = pg.inner_text('.list-total')
        ok(f'{int(n):,} سجل' in t and money(s) in t, f'{page}: all years = {n} records, total {money(s)} ({t!r})')
    list_total('oil', 'oil_changes', 'change_date', 'oil list (1,6xx)')
    list_total('maintenance', 'maintenance', 'maint_date', 'maintenance list')
    list_total('food', 'food_supplies', 'supply_date', 'food list (1,2xx)')
    list_total('housing', 'housing_maintenance', 'work_date', 'housing list')
    list_total('accidents', 'accidents', 'occurred_at', 'accidents list')

    # reports: all-period expenses must equal the database total
    timed('reports: expenses', lambda: (pg.goto(BASE + '#/reports'), pg.wait_for_selector('.rep-head')))
    timed('reports: all periods', lambda: (pg.click('[data-q=all]'), pg.wait_for_timeout(200), pg.wait_for_selector('.rep-stats')))
    pg.wait_for_timeout(1500)
    grand = sql('select sum(total) from v_expenses')
    ok(money(grand) in pg.inner_text('.rep-stats'), f'reports: all-period expenses {money(grand)} = database (4,121 operations, > 1000-row cap)')
    for y in sorted({int(x) for x in sql("select distinct extract(year from expense_date) from v_expenses").split()}):
        ys = sql(f'select sum(total) from v_expenses where extract(year from expense_date) = {y}')
        ok(money(ys) in pg.inner_text('.rep-stats'), f'reports: year {y} = {money(ys)}')
    pg.select_option('select[name=report]', 'cost_company'); pg.wait_for_timeout(1500)
    for co, name in [('11111111-1111-1111-1111-111111111111', 'روى'), ('22222222-2222-2222-2222-222222222222', 'الفهاد')]:
        s = sql(f"select sum(total) from v_expenses where site_company_id = '{co}'")
        ok(money(s) in pg.inner_text('.rep-table'), f'cost per company {name}: {money(s)}')
    pg.select_option('select[name=report]', 'documents'); pg.wait_for_timeout(1500)
    exp = sql("select count(*) from v_document_status where status='expired'")
    ok(f'{int(exp):,}' in pg.inner_text('.rep-stats'), f'documents report: {exp} expired (all rows)')
    pg.screenshot(path=f'{OUT}/bulk-report.png', full_page=False)

    # dashboard for a past month with data
    pg.goto(BASE + '#/dashboard'); pg.wait_for_selector('.panels')

    # Excel export sees every row
    pg.goto(BASE + '#/excel'); pg.wait_for_selector('.exp-grid')
    for key, table, sheet in [('oil', 'oil_changes', 'تغيير الزيت'), ('food', 'food_supplies', 'توريد الطعام'), ('vehicles', 'vehicles', 'السيارات')]:
        t0 = time.time()
        with pg.expect_download(timeout=120000) as dl: pg.click(f'[data-exp={key}]')
        path = f'{OUT}/bulk-{key}.xlsx'; dl.value.save_as(path); timings.append((f'excel export {key}', time.time() - t0))
        ws = openpyxl.load_workbook(path)[sheet]
        n = int(sql(f'select count(*) from {table} where deleted_at is null'))
        ok(ws.max_row - 1 == n, f'excel {key}: {ws.max_row - 1} rows = database {n}')

    # backup sees every row
    pg.goto(BASE + '#/backup'); pg.wait_for_selector('[data-act=download]')
    t0 = time.time()
    with pg.expect_download(timeout=180000) as dl: pg.click('[data-act=download]')
    bp = f'{OUT}/bulk-backup.json'; dl.value.save_as(bp); timings.append(('full backup download', time.time() - t0))
    bk = json.load(open(bp))
    tables = bk['backup']['tables'] if 'backup' in bk else bk.get('tables', bk)
    for t in ['vehicles', 'documents', 'oil_changes', 'food_supplies', 'audit_log', 'odometer_readings']:
        n = int(sql(f'select count(*) from {t}')) - (1 if t == 'audit_log' else 0)  # the backup itself records «last backup» afterwards
        ok(len(tables[t]) == n, f'backup {t}: {len(tables[t])} = {n}')
    # search across 1,105 vehicles
    timed('search', lambda: (pg.fill('form.search input', '11050'), pg.press('form.search input', 'Enter'), pg.wait_for_selector('td.plate')))
    ok('ت ج ر 11050' in pg.inner_text('#page'), 'search finds vehicle 1050 of 1100')
    b.close()

print(f'PASS {len(passed)}')
for x in passed: print('  ok  ', x)
print('timings (4x CPU slowdown):')
for l, t in timings: print(f'  {t:6.2f}s  {l}')
print(f'FAIL {len(problems)}')
for x in problems: print('  FAIL', x)
