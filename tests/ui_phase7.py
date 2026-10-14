"""Phase 7 browser test: Excel export (validated with openpyxl), Excel import round-trip with errors, reports.
Usage: python3 -I ui_phase7.py OUTDIR   (run reset_ui_db.sh + start_local.sh first)"""
import sys, os, subprocess, datetime, shutil
from playwright.sync_api import sync_playwright
import openpyxl

OUT = sys.argv[1]; os.makedirs(OUT, exist_ok=True)
BASE = 'http://localhost:8787/'
problems, passed = [], []
def ok(c, m): (passed if c else problems).append(m)
PSQL = ['psql', '-h', '/tmp', '-p', '5433', '-U', 'postgres', '-X', '-q', '-At', '-d', 'fleet_ui']
def sql(q): return subprocess.run(PSQL + ['-c', q], capture_output=True, text=True).stdout.strip()

def login(pg, email):
    pg.goto(BASE); pg.wait_for_selector('form.login-card')
    pg.fill('input[name=email]', email); pg.fill('input[name=password]', 'test1234'); pg.click('button[type=submit]')
    pg.wait_for_selector('.panels')

def new_page(b, w=390):
    c = b.new_context(viewport={'width': w, 'height': 844}, device_scale_factor=2, locale='ar-SA', timezone_id='Asia/Riyadh', accept_downloads=True)
    g = c.new_page(); g.route('**/fonts.googleapis.com/**', lambda r: r.abort())
    g.on('pageerror', lambda e: problems.append(f'page error: {e}'))
    return c, g

def download(pg, selector, name):
    with pg.expect_download() as dl: pg.click(selector)
    path = f'{OUT}/{name}'; dl.value.save_as(path); return path, dl.value.suggested_filename

def headers(ws): return [c.value for c in ws[1]]
TODAY = datetime.date.fromisoformat(sql("select app_today()"))
before = {t: sql(f'select count(*) from {t}') for t in ['vehicles', 'employees', 'documents', 'oil_changes', 'maintenance', 'accidents', 'attachments', 'odometer_readings']}

with sync_playwright() as p:
    b = p.chromium.launch()
    ctx, pg = new_page(b, 1280)
    login(pg, 'admin@test.local')

    # ======================= EXPORT =======================
    pg.goto(BASE + '#/excel'); pg.wait_for_selector('.exp-grid')
    ok(pg.locator('.exp-grid [data-exp]').count() == 13, 'excel page offers 13 export buttons + all')
    pg.screenshot(path=f'{OUT}/p7-excel.png', full_page=True)
    path, fname = download(pg, '[data-exp=vehicles]', 'vehicles.xlsx')
    ok(fname.endswith('.xlsx') and 'السيارات' in fname, f'vehicles export file name ({fname})')
    wb = openpyxl.load_workbook(path); ws = wb.active
    h = headers(ws)
    ok(ws.title == 'السيارات' and ws.sheet_view.rightToLeft, 'sheet named السيارات, right-to-left')
    ok(ws.freeze_panes == 'A2' and ws.auto_filter.ref, 'header row frozen + autofilter')
    ok(ws.max_row == 6 and h[0] == 'الرقم الداخلي' and 'انتهاء الاستمارة' in h and 'انتهاء الفحص الدوري' in h, 'vehicles: 5 rows, Arabic headers incl. expiry columns')
    codes = [ws.cell(r, 1).value for r in range(2, 7)]
    ok(codes == ['VEH-0001', 'VEH-0002', 'VEH-0003', 'VEH-0004', 'VEH-0005'], 'rows keyed by internal code')
    reg = ws.cell(2, h.index('انتهاء الاستمارة') + 1)
    ok(isinstance(reg.value, datetime.datetime) and reg.value.date() == TODAY + datetime.timedelta(days=210) and reg.number_format == 'yyyy/mm/dd',
       f'expiry is a real date cell yyyy/mm/dd ({reg.value}, {reg.number_format})')
    odo = ws.cell(2, h.index('قراءة العداد') + 1).value
    ok(odo == 91300 and isinstance(odo, int), 'odometer exported as number 91300')
    ok(ws.cell(2, h.index('الشركة المالكة') + 1).value == 'شركة روى الرمال الزراعية', 'company names in Arabic')
    ok(ws.cell(5, h.index('الحالة التشغيلية') + 1).value == 'متضررة بحادث', 'enum exported as Arabic label')

    path_e, _ = download(pg, '[data-exp=employees]', 'employees.xlsx')
    we = openpyxl.load_workbook(path_e).active
    ok(we.max_row == 6 and 'رقم الإقامة' in headers(we), 'employees export (admin sees sensitive columns)')

    path_m, _ = download(pg, '[data-exp=maintenance]', 'maint.xlsx')
    wm = openpyxl.load_workbook(path_m)
    ok(wm.sheetnames == ['الصيانة', 'بنود الصيانة'], f'maintenance export has items sheet ({wm.sheetnames})')
    mh = headers(wm['الصيانة']); tot = wm['الصيانة'].cell(2, mh.index('الإجمالي') + 1)
    ok(tot.value == 1380 and '0.00' in tot.number_format, f'maintenance total 1380 as money cell ({tot.value}, {tot.number_format})')
    ok(wm['بنود الصيانة'].max_row == 3, 'two maintenance item rows')

    path_a, _ = download(pg, '[data-exp=all]', 'all.xlsx')
    wa = openpyxl.load_workbook(path_a)
    ok(len(wa.sheetnames) == 14 and all(wa[s].sheet_view.rightToLeft for s in wa.sheetnames), f'«all» export: 14 RTL sheets ({len(wa.sheetnames)})')
    acc = wa['الحوادث']; ah = headers(acc); c = acc.cell(2, [i for i, x in enumerate(ah) if 'تاريخ' in (x or '')][0] + 1)
    ok(isinstance(c.value, datetime.datetime) and 'hh:mm' in c.number_format, f'accident date-time cell ({c.number_format})')
    ok('آخر' in pg.inner_text('[data-exp-msg]') or 'تم التصدير' in pg.inner_text('[data-exp-msg]'), 'export confirmation message')

    # ======================= IMPORT: vehicles =======================
    wb = openpyxl.load_workbook(path); ws = wb.active; col = {x: i + 1 for i, x in enumerate(headers(ws))}
    ws.cell(6, col['اللون']).value = 'فضي'                                              # VEH-0005 update
    ws.cell(2, col['انتهاء الاستمارة']).value = datetime.datetime(2028, 3, 1)           # VEH-0001 new registration doc
    ws.cell(3, col['قراءة العداد']).value = 61000                                       # VEH-0002 odometer lower -> error
    ws.cell(4, col['انتهاء الفحص الدوري']).value = None                                  # empty cell must not wipe a doc
    nr = 7
    def add(**kv):
        global nr
        for k, v in kv.items(): ws.cell(nr, col[k]).value = v
        nr += 1
    add(**{'رقم اللوحة': 'ن هـ ي 4321', 'الماركة': 'فورد', 'نوع السيارة': 'بيك أب', 'سنة الصنع': 2024, 'الشركة المالكة': 'شركة عبد الرحمن الفهاد',
           'الموقع': 'مشتل الغملاس', 'الحالة التشغيلية': 'تعمل', 'قراءة العداد': 1500, 'انتهاء الاستمارة': '2027/05/20', 'انتهاء الفحص الدوري': '٢٠٢٧/٠١/١٠'})
    add(**{'رقم اللوحة': 'ب ب ب 1111', 'الشركة المالكة': 'شركة روى الرمال الزراعية', 'انتهاء الاستمارة': '2026/02/31'})       # invalid date
    add(**{'رقم اللوحة': 'أ ب ج 1234', 'الشركة المالكة': 'شركة روى الرمال الزراعية'})                                     # duplicate plate
    add(**{'رقم اللوحة': 'ج ج ج 2222', 'الشركة المالكة': 'شركة غير موجودة'})                                            # unknown company
    add(**{'رقم اللوحة': 'د د د 3333', 'الشركة المالكة': 'شركة روى الرمال الزراعية', 'الحالة التشغيلية': 'طائرة'})       # bad enum
    add(**{'الرقم الداخلي': 'VEH-9999', 'رقم اللوحة': 'هـ هـ هـ 4444', 'الشركة المالكة': 'شركة روى الرمال الزراعية'})   # unknown code
    add(**{'رقم اللوحة': None, 'الشركة المالكة': 'شركة روى الرمال الزراعية', 'الماركة': 'تويوتا'})                       # missing plate
    imp = f'{OUT}/vehicles-edited.xlsx'; wb.save(imp)

    with pg.expect_file_chooser() as fc: pg.click('[data-imp=vehicles]')
    fc.value.set_files(imp); pg.wait_for_selector('.imp-chips')
    chips = {k: int(pg.inner_text(f'.imp-chips [data-f={k}] b')) for k in ['new', 'update', 'same', 'error', 'all']}
    ok(chips == {'new': 1, 'update': 2, 'same': 2, 'error': 7, 'all': 12}, f'preview counts new/update/same/error ({chips})')
    errs = pg.inner_text('.imp-rows')
    for frag, what in [('غير موجود', 'invalid 2026/02/31'), ('مسجل لسيارة أخرى', 'duplicate plate'), ('غير مسجلة', 'unknown company'),
                       ('غير معروفة', 'bad Arabic status'), ('VEH-9999', 'unknown code'), ('رقم اللوحة» مطلوب', 'missing plate'), ('أقل من آخر قراءة', 'odometer decrease')]:
        ok(frag in errs, f'error shown: {what}')
    pg.screenshot(path=f'{OUT}/p7-import-errors.png', full_page=True)
    pg.click('.imp-chips [data-f=update]'); t = pg.inner_text('.imp-rows')
    ok('فضي' in t and '2028/03/01' in t and 'VEH-0005' in t, 'update preview shows old → new (color, registration date)')
    pg.screenshot(path=f'{OUT}/p7-import-updates.png', full_page=True)
    ok(sql("select count(*) from vehicles") == before['vehicles'], 'nothing saved before approval')
    pg.click('[data-apply]'); pg.wait_for_selector('.modal'); pg.click('.modal [data-yes]')
    pg.wait_for_selector('.imp-box .warn-box')
    ok('إضافة 1' in pg.inner_text('.imp-box').replace('\n', ' ') and 'تعديل 2' in pg.inner_text('.imp-box').replace('\n', ' '), 'import result: 1 added, 2 updated')
    ok(sql("select count(*) from vehicles") == str(int(before['vehicles']) + 1), 'one vehicle added, none deleted')
    nv = sql("select code||'|'||make||'|'||manufacture_year||'|'||odometer||'|'||op_status from vehicles where plate_no='ن هـ ي 4321'")
    ok(nv == 'VEH-0006|فورد|2024|1500|working', f'new vehicle saved with generated code ({nv})')
    ok(sql("select string_agg(doc_type||'='||expiry_date, ',' order by doc_type::text) from documents d join vehicles v on v.id=d.vehicle_id where v.plate_no='ن هـ ي 4321'")
       == 'periodic_inspection=2027-01-10,registration=2027-05-20', 'new vehicle documents (incl. Arabic digits date)')
    ok(sql("select color from vehicles where code='VEH-0005'") == 'فضي', 'VEH-0005 color updated')
    ok(sql("select count(*) from documents where vehicle_id='b0000000-0000-0000-0000-000000000001' and doc_type='registration'") == '2', 'old registration kept + new one added (history)')
    ok(sql("select count(*) from documents where vehicle_id='b0000000-0000-0000-0000-000000000003' and doc_type='periodic_inspection' and deleted_at is null") == '1', 'empty cell did not wipe a document')
    ok(sql("select odometer from vehicles where code='VEH-0002'") == '61400', 'rejected odometer decrease not applied')
    ok(sql("select summary->>'inserted'||'/'||(summary->>'updated')||'/'||(summary->>'rejected') from import_batches where target_table='vehicles'") == '1/2/7', 'import batch logged')
    ok(sql("select count(*) from audit_log where table_name='vehicles' and action='UPDATE' and record_id='b0000000-0000-0000-0000-000000000005'") >= '1', 'import changes in audit log')
    for t in ['oil_changes', 'maintenance', 'accidents', 'attachments']:
        ok(sql(f'select count(*) from {t}') == before[t], f'{t} untouched by import')

    # re-import same file → nothing new
    path2, _ = download(pg, '[data-exp=vehicles]', 'vehicles2.xlsx')
    with pg.expect_file_chooser() as fc: pg.click('[data-imp=vehicles]')
    fc.value.set_files(path2); pg.wait_for_selector('.imp-chips')
    ok(pg.inner_text('.imp-chips [data-f=same] b') == '6' and pg.inner_text('.imp-chips [data-f=update] b') == '0', 're-importing a fresh export = 6 unchanged, 0 updates')

    # ======================= IMPORT: employees =======================
    we_wb = openpyxl.load_workbook(path_e); we = we_wb.active; ec = {x: i + 1 for i, x in enumerate(headers(we))}
    we.cell(6, ec['المسمى الوظيفي']).value = 'مشرف عام'
    we.cell(2, ec['رقم الجوال']).value = '0500000001'
    we.cell(7, ec['الاسم الكامل']).value = 'موظف مستورد تجريبي'; we.cell(7, ec['الموقع']).value = 'مزرعة ارم'; we.cell(7, ec['رقم الإقامة']).value = '2000000001'
    we.cell(7, ec['انتهاء الإقامة']).value = '2027/12/31'
    we.cell(8, ec['الاسم الكامل']).value = 'موظف خطأ'; we.cell(8, ec['رقم الإقامة']).value = '123'
    we.cell(9, ec['الاسم الكامل']).value = 'موظف خطأ 2'; we.cell(9, ec['حالة الموظف']).value = 'غير معروف'
    impe = f'{OUT}/employees-edited.xlsx'; we_wb.save(impe)
    with pg.expect_file_chooser() as fc: pg.click('[data-imp=employees]')
    fc.value.set_files(impe); pg.wait_for_selector('.imp-chips')
    ce = {k: int(pg.inner_text(f'.imp-chips [data-f={k}] b')) for k in ['new', 'update', 'same', 'error']}
    ok(ce == {'new': 1, 'update': 2, 'same': 3, 'error': 2}, f'employees preview ({ce})')
    ok('10 أرقام' in pg.inner_text('.imp-rows'), 'iqama must be 10 digits')
    pg.click('[data-apply]'); pg.wait_for_selector('.modal'); pg.click('.modal [data-yes]'); pg.wait_for_selector('.imp-box .warn-box')
    ok(sql("select e.company_id||'|'||p.iqama_no from employees e join employee_private p on p.employee_id=e.id where e.full_name='موظف مستورد تجريبي'")
       == '11111111-1111-1111-1111-111111111111|2000000001', 'new employee: company from site, iqama in private table')
    ok(sql("select mobile from employee_private where employee_id='e0000000-0000-0000-0000-000000000001'") == '0500000001', 'mobile updated in private table')
    ok(sql("select job_title from employees where id='e0000000-0000-0000-0000-000000000005'") == 'مشرف عام', 'job title updated')

    # ======================= REPORTS =======================
    pg.goto(BASE + '#/reports'); pg.wait_for_selector('.rep-head')
    pg.click('[data-q=all]'); pg.wait_for_selector('.rep-stats'); pg.wait_for_timeout(300)
    stats = pg.inner_text('.rep-stats')
    ok('11,305.00' in stats, f'expenses total 11,305.00 SAR for all periods ({stats[:80]!r})')
    ok(pg.locator('.legend-row span').count() == 5 and pg.locator('svg.stack rect.s1').count() >= 1, 'stacked monthly chart with 5-category legend')
    ok(pg.locator('.rep-table tbody tr').count() >= 1, 'expenses table view present')
    pg.hover('svg.stack g.col >> nth=0'); pg.wait_for_timeout(150)
    ok(not pg.locator('#tip').is_hidden() and 'ر.س' in pg.inner_text('#tip'), 'chart tooltip on hover')
    pg.screenshot(path=f'{OUT}/p7-report-expenses.png', full_page=True)
    with pg.expect_download() as dl: pg.click('[data-x]')
    rp = f'{OUT}/report.xlsx'; dl.value.save_as(rp); rw = openpyxl.load_workbook(rp).active
    rh = headers(rw)
    ok(rh[0] == 'الشهر' and rh[-1] == 'الإجمالي' and abs(sum(rw.cell(r, len(rh)).value or 0 for r in range(2, rw.max_row + 1)) - 11305) < 0.01, 'report Excel export sums to 11,305')

    pg.select_option('select[name=report]', 'cost_site'); pg.wait_for_timeout(500)
    rows = pg.inner_text('.rep-table tbody')
    ok('مشتل الغملاس' in rows and pg.locator('.hb').count() == 4, 'cost per site: 4 sites')
    pg.select_option('select[name=company]', '22222222-2222-2222-2222-222222222222'); pg.wait_for_timeout(500)
    # company 2 (nursery): accident 3300 + housing 690 + food 782 = 4772
    ok('4,772.00' in pg.inner_text('.rep-stats') and pg.locator('.hb').count() == 1, 'company filter: nursery total 4,772.00')
    pg.select_option('select[name=report]', 'cost_company'); pg.wait_for_timeout(500)
    ok('4,772.00' in pg.inner_text('.rep-table'), 'cost per company matches')
    pg.select_option('select[name=company]', ''); pg.wait_for_timeout(400)
    pg.select_option('select[name=report]', 'cost_vehicle'); pg.wait_for_timeout(500)
    ok('3,300.00' in pg.inner_text('.rep-table') and 'ط ع ق 1357' in pg.inner_text('.rep-table'), 'cost per vehicle: accident vehicle 3,300')
    pg.hover('.hb >> nth=0'); pg.wait_for_timeout(100)
    ok('3,300.00' in pg.inner_text('#tip'), 'bar tooltip shows value')
    pg.select_option('select[name=report]', 'documents'); pg.wait_for_timeout(500)
    d = pg.inner_text('.rep-stats')
    ok(sql("select count(*) from v_document_status where status='expired'") in d, 'documents report expired count matches DB')
    pg.select_option('select[name=report]', 'vehicles_status'); pg.wait_for_timeout(500)
    ok('6' in pg.inner_text('.rep-stats') and pg.locator('.hb').count() == 4, 'vehicles by status (6 vehicles, 4 statuses)')

    # date filter: this month only, then invalid date
    pg.select_option('select[name=report]', 'food'); pg.wait_for_timeout(400)
    pg.fill('input[name=from]', '2026/02/31'); pg.dispatch_event('input[name=from]', 'change'); pg.wait_for_timeout(300)
    ok(pg.locator('.toast').count() >= 1, 'invalid report date rejected with message')
    fr = (TODAY - datetime.timedelta(days=3)).strftime('%Y/%m/%d')
    pg.fill('input[name=from]', fr); pg.dispatch_event('input[name=from]', 'change'); pg.wait_for_timeout(500)
    ok('1,955.00' in pg.inner_text('.rep-stats'), 'food from today-3: only 1,955.00')

    # every report renders without error
    vals = pg.eval_on_selector_all('select[name=report] option', 'os => os.map(o => o.value)')
    pg.click('[data-q=all]'); pg.wait_for_timeout(300)
    bad = []
    for v in vals:
        pg.select_option('select[name=report]', v); pg.wait_for_timeout(350)
        if pg.locator('.report .bad-box').count() or not pg.locator('.rep-head').count(): bad.append(v)
    ok(len(vals) == 15 and not bad, f'all 15 reports render ({bad})')

    pg.select_option('select[name=report]', 'expenses'); pg.wait_for_timeout(400)
    pg.emulate_media(media='print'); pg.screenshot(path=f'{OUT}/p7-report-print.png', full_page=True)
    ok(pg.locator('.rep-filters').is_hidden() and not pg.locator('.print-only').is_hidden(), 'print view hides filters, shows footer')
    pg.emulate_media(media='screen')

    # mobile
    c3, m = new_page(b, 390); login(m, 'admin@test.local')
    m.goto(BASE + '#/reports'); m.wait_for_selector('.rep-head'); m.click('[data-q=all]'); m.wait_for_timeout(400)
    ok(m.evaluate('document.documentElement.scrollWidth <= window.innerWidth + 1'), 'mobile report has no horizontal page scroll')
    m.screenshot(path=f'{OUT}/p7-report-mobile.png', full_page=True)
    m.select_option('select[name=report]', 'cost_site'); m.wait_for_timeout(400)
    m.screenshot(path=f'{OUT}/p7-report-mobile-site.png', full_page=True)
    m.goto(BASE + '#/excel'); m.wait_for_selector('.exp-grid'); m.screenshot(path=f'{OUT}/p7-excel-mobile.png', full_page=True)
    c3.close()

    # viewer: no export/import, report has no Excel button
    c2, v = new_page(b, 390); login(v, 'viewer@test.local')
    v.goto(BASE + '#/excel'); v.wait_for_selector('.panel')
    ok(v.locator('[data-exp]').count() == 0 and v.locator('[data-imp]').count() == 0, 'viewer: no export / import buttons')
    v.goto(BASE + '#/reports'); v.wait_for_selector('.rep-head')
    ok(v.locator('[data-x]').count() == 0 and v.locator('[data-p]').count() == 1, 'viewer: report printable but no Excel export')
    c2.close()
    b.close()

print(f'PASS {len(passed)}')
for x in passed: print('  ok  ', x)
print(f'FAIL {len(problems)}')
for x in problems: print('  FAIL', x)
