"""Maintenance requests (replaces periodic maintenance screen): create, auto-fill, validation, approval flow,
lock, attachments, print/PDF, vehicle file link, search, alerts, permissions, old maintenance kept, Excel.
Usage: python3 -I ui_requests.py OUTDIR   (fleet_ui reset with ui_seed.sql)"""
import sys, os, subprocess
from playwright.sync_api import sync_playwright
import openpyxl
import datetime as _dt, zoneinfo as _zi
_T0 = _dt.datetime.now(_zi.ZoneInfo('Asia/Riyadh')).date()
def _d(n=0, sep='/'): return (_T0 + _dt.timedelta(days=n)).strftime(f'%Y{sep}%m{sep}%d')

OUT = sys.argv[1]; os.makedirs(OUT, exist_ok=True)
HERE = os.path.dirname(os.path.abspath(__file__)); FX = os.path.join(HERE, 'fixtures')
BASE = 'http://localhost:8787/'
V1, V5 = 'b0000000-0000-0000-0000-000000000001', 'b0000000-0000-0000-0000-000000000005'
problems, passed = [], []
def ok(c, m): (passed if c else problems).append(m)
PSQL = ['psql', '-h', '/tmp', '-p', '5433', '-U', 'postgres', '-X', '-q', '-At', '-d', 'fleet_ui']
def sql(q): return subprocess.run(PSQL + ['-c', q], capture_output=True, text=True).stdout.strip()

def ctx(b, w=390):
    c = b.new_context(viewport={'width': w, 'height': 900}, device_scale_factor=2, locale='ar-SA', timezone_id='Asia/Riyadh', accept_downloads=True)
    g = c.new_page(); g.route('**/fonts.googleapis.com/**', lambda r: r.abort())
    g.on('pageerror', lambda e: problems.append(f'page error: {e}'))
    g.add_init_script('window.__prints = 0; window.print = () => { window.__prints++; };')
    return c, g
def login(g, email='admin@test.local'):
    g.goto(BASE); g.wait_for_selector('form.login-card')
    g.fill('input[name=email]', email); g.fill('input[name=password]', 'test1234'); g.click('button[type=submit]'); g.wait_for_selector('.panels')
def fill(g, vals):
    for n, v in vals.items():
        el = g.locator(f'.modal [name="{n}"]').last
        if el.evaluate('e => e.tagName') == 'SELECT': el.select_option(value=v)
        elif el.get_attribute('type') == 'checkbox': el.check() if v else el.uncheck()
        else: el.fill(v)
def val(g, n): return g.locator(f'.modal [name="{n}"]').last.input_value()

with sync_playwright() as p:
    b = p.chromium.launch()
    c, g = ctx(b, 1280); login(g)
    ok(g.locator('.sidebar [data-nav=requests]').inner_text().strip() == 'طلب صيانة' and g.locator('.sidebar [data-nav=maintenance]').count() == 0,
       'sidebar: «الصيانة الدورية» replaced by «طلب صيانة»')

    # ---------- create from the list with a registered vehicle
    g.click('.sidebar [data-nav=requests]'); g.wait_for_selector('.list-body')
    ok('لا توجد طلبات' in g.inner_text('.list-body'), 'empty list message')
    g.click('[data-list-add]'); g.wait_for_selector('.modal [name=fault]')
    ok(_d() in g.inner_text('.modal .mr-head') and 'يُنشأ تلقائياً' in g.inner_text('.modal .mr-head'), 'form shows automatic date and number')
    fill(g, {'vehicle_id': V1}); g.wait_for_timeout(200)
    ok(val(g, 'plate_no') == 'أ ب ج 1234' and 'هايلكس' in val(g, 'model') and val(g, 'driver_name') == 'سائق تجريبي 1' and val(g, 'meter_reading') == '91300',
       f"vehicle auto-fills plate, model, user, odometer ({val(g, 'plate_no')}, {val(g, 'model')}, {val(g, 'driver_name')}, {val(g, 'meter_reading')})")
    ok(g.locator('.modal [name=site_id]').input_value() != '', 'project auto-selected from the vehicle site')
    g.click('.modal [data-save]'); g.wait_for_timeout(200)
    ok('العطل' in g.inner_text('.modal .form-error'), 'fault is required')
    g.screenshot(path=f'{OUT}/r-form-desktop.png', full_page=True)
    fill(g, {'project_text': 'مشروع تمديد الري', 'model': 'هايلكس 2022 غمارتين', 'fault': 'صوت طقطقة في المحرك عند التشغيل\nوارتفاع الحرارة',
             'fault_cause': 'احتمال تلف سير المحرك', 'technician_opinion': 'فحص السير والمضخة', 'site_notes': 'السيارة مطلوبة للعمل يومياً',
             'last_change_date': '2026/09/01', 'meter_reading': '91450', 'send': True})
    g.click('.modal [data-save]'); g.wait_for_selector('.mr-sheet')
    ok('#/requests/' in g.url, 'after save, the request page opens')
    sheet = g.inner_text('.mr-sheet')
    for frag, what in [('طلب صيانة', 'title'), ('MR-000001', 'number'), (_d(), 'date'), ('مزرعة الغملاس', 'project'), ('مشروع تمديد الري', 'project text'),
                       ('أ ب ج 1234', 'plate'), ('هايلكس 2022 غمارتين', 'model typed over'), ('سائق تجريبي 1', 'user'), ('91,450 كم', 'meter'),
                       ('صوت طقطقة في المحرك', 'fault'), ('احتمال تلف سير المحرك', 'cause'), ('فحص السير والمضخة', 'technician'), ('مطلوبة للعمل يومياً', 'site notes'),
                       ('2026/09/01', 'last change date'), ('بانتظار اعتماد المشرف', 'status'), ('مشرف الصيانة', 'supervisor box'), ('الاعتماد', 'approval box')]:
        ok(frag in sheet, f'sheet shows {what}')
    ok(sql("select status||'|'||meter_reading||'|'||(vehicle_id is not null)||'|'||(driver_id is not null) from maintenance_requests where request_no='MR-000001'") == 'pending_supervisor|91450|true|true',
       'saved linked to vehicle and driver, sent for approval')
    ok(sql(f"select odometer from vehicles where id='{V1}'") == '91450', 'vehicle odometer raised to the request reading')
    g.screenshot(path=f'{OUT}/r-page-desktop.png', full_page=True)

    # ---------- attachments on the request
    with g.expect_file_chooser() as fc: g.locator('.att-box [data-up-file]').first.click()
    fc.value.set_files([os.path.join(FX, 'invoice.jpg')]); g.wait_for_selector('.modal [name=category]')
    cats = g.locator('.modal [name=category] option').all_inner_texts()
    ok('صورة العطل' in cats and 'تقرير فني' in cats and 'فاتورة صيانة' in cats, f'request attachment types: fault photo, report, invoice ({cats})')
    g.locator('.modal [name=category]').select_option('fault_photo'); g.click('.modal [data-save]')
    g.wait_for_selector('.modal [name=category]', state='detached', timeout=30000); g.wait_for_timeout(500)
    ok(sql("select count(*) from attachments where entity_type='maintenance_request' and category='fault_photo'") == '1', 'fault photo attached to the request')

    # ---------- print / PDF
    g.click('[data-act=print]'); g.wait_for_timeout(100)
    g.click('[data-act=pdf]'); g.wait_for_timeout(700)
    ok(g.evaluate('window.__prints') == 2, 'print and PDF buttons open the print dialog')
    g.emulate_media(media='print'); g.screenshot(path=f'{OUT}/r-print.png', full_page=True)
    ok(g.locator('.mr-sheet').is_visible() and g.locator('.mr-bar').is_hidden() and g.locator('.att-box').is_hidden() and g.locator('.sidebar').is_hidden(),
       'print shows only the form (no buttons, menus, attachments)')
    g.emulate_media(media='screen')
    pdf = g.pdf(path=f'{OUT}/r-request.pdf', format='A4') if False else None

    # ---------- approval
    ok(g.locator('[data-act=decide]').count() == 1, 'admin sees «اعتماد أو رفض»')
    g.click('[data-act=decide]'); g.wait_for_selector('.modal [name=decision]')
    fill(g, {'supervisor_name': 'م. خالد المشرف', 'approver_name': 'أ. سعد المدير', 'approval_notes': 'يُنفّذ في ورشة الموقع'})
    g.click('.modal [data-save]'); g.wait_for_selector('.modal', state='detached'); g.wait_for_selector('.mr-sheet:has-text("أ. سعد المدير")'); g.wait_for_selector('[data-act=start]')
    t = g.inner_text('.mr-sheet')
    ok('م. خالد المشرف' in t and 'أ. سعد المدير' in t and 'معتمد' in t, 'sheet shows supervisor, approver and «معتمد»')
    ok(g.locator('[data-act=start]').count() == 1 and g.locator('[data-act=decide]').count() == 0, 'after approval: «بدء الصيانة» shown, approval hidden')
    g.click('[data-act=edit]'); g.wait_for_selector('.modal [name=technician_opinion]')
    ok(g.locator('.modal [name=fault]').count() == 0 and g.locator('.modal [name=technician_opinion]').count() == 1, f"approved request: only technician opinion and notes editable ({g.eval_on_selector_all('.modal [name]', 'es => es.map(e => e.name)')})")
    fill(g, {'technician_opinion': 'تم تغيير سير المحرك وفحص المضخة'}); g.click('.modal [data-save]'); g.wait_for_selector('.modal', state='detached'); g.wait_for_timeout(400)
    ok('تم تغيير سير المحرك' in g.inner_text('.mr-sheet'), 'technician opinion updated after approval')
    g.click('[data-act=start]'); g.wait_for_timeout(600)
    ok('قيد الصيانة' in g.inner_text('.mr-track .tab-head'), 'status → قيد الصيانة')
    g.click('[data-act=done]'); g.wait_for_selector('.modal [data-yes]'); g.click('.modal [data-yes]'); g.wait_for_timeout(700)
    ok(sql("select status||'|'||(started_at is not null)||'|'||(completed_at is not null) from maintenance_requests where request_no='MR-000001'") == 'completed|true|true', 'status → مكتمل with times')
    ok(g.locator('[data-act=invoice]').count() == 1, '«تسجيل فاتورة الصيانة» offered after completion')
    g.click('[data-act=invoice]'); g.wait_for_selector('.modal [name=maint_type]')
    fill(g, {'maint_type': 'إصلاح محرك', 'workshop': 'ورشة الموقع'})
    g.locator('.modal [name=it_desc]').first.fill('سير محرك'); g.locator('.modal [name=it_price]').first.fill('350')
    g.click('.modal [data-save]'); g.wait_for_selector('.modal', state='detached'); g.wait_for_timeout(700)
    ok(sql("select m.maint_type||'|'||m.subtotal from maintenance_requests r join maintenance m on m.id = r.maintenance_id where r.request_no='MR-000001'") == 'إصلاح محرك|350.00',
       'maintenance invoice recorded and linked to the request (costs stay in reports)')
    g.screenshot(path=f'{OUT}/r-completed.png', full_page=True)

    # ---------- equipment not registered + rejection
    g.click('[data-act=new]'); g.wait_for_selector('.modal [name=fault]')
    fill(g, {'vehicle_id': V5}); g.wait_for_timeout(150)
    fill(g, {'vehicle_id': ''}); g.wait_for_timeout(100)
    fill(g, {'plate_no': '', 'model': '', 'driver_name': '', 'meter_reading': ''})
    fill(g, {'fault': 'لا يعمل'}); g.click('.modal [data-save]'); g.wait_for_timeout(300)
    ok('المعدة' in g.inner_text('.modal .form-error'), 'needs a vehicle, equipment name or plate')
    fill(g, {'equipment_text': 'مولد كهرباء 50 ك.ف', 'plate_no': 'بدون', 'model': 'كاتربلر', 'driver_name': 'عامل الموقع', 'meter_reading': '1200', 'meter_unit': 'hours',
             'project_text': 'سكن العمال', 'send': False})
    g.click('.modal [data-save]'); g.wait_for_selector('.mr-sheet'); g.wait_for_timeout(300)
    t = g.inner_text('.mr-sheet')
    ok('MR-000002' in t and 'مولد كهرباء 50 ك.ف' in t and '1,200 ساعة' in t and 'جديد' in t, 'equipment request with manual data, hours meter, status «جديد»')
    g.click('[data-act=send]'); g.wait_for_timeout(500)
    g.click('[data-act=decide]'); g.wait_for_selector('.modal [name=decision]')
    fill(g, {'decision': 'rejected', 'approval_notes': ''}); g.click('.modal [data-save]'); g.wait_for_timeout(300)
    ok('سبب' in g.inner_text('.modal .form-error'), 'rejection needs a reason')
    fill(g, {'approval_notes': 'يُستبدل المولد بدلاً من إصلاحه'}); g.click('.modal [data-save]'); g.wait_for_selector('.modal', state='detached'); g.wait_for_timeout(500)
    t = g.inner_text('.mr-sheet')
    ok('مرفوض' in t and 'سبب الرفض' in t and 'يُستبدل المولد' in t, 'rejected request shows status and reason')

    # ---------- vehicle file: requests tab + new request prefilled from the vehicle
    g.goto(BASE + f'#/vehicles/{V1}?tab=maint'); g.wait_for_selector('.vhead')
    pnl = g.inner_text('[data-panel=maint]')
    ok('طلبات الصيانة' in g.inner_text('.tabs') and 'MR-000001' in pnl and _d() in pnl and 'مكتمل' in pnl, 'vehicle file tab «طلبات الصيانة» shows number, date, status')
    ok('تاريخ آخر تغيير' in pnl and '2026/09/01' in pnl, 'vehicle tab shows last change date')
    ok('فواتير الصيانة السابقة' in pnl and 'MNT-' in pnl, 'previous maintenance invoices kept in the vehicle file')
    g.click('[data-panel=maint] [data-act=mreq]'); g.wait_for_selector('.modal [name=fault]')
    ok(val(g, 'plate_no') == 'أ ب ج 1234', 'new request from the vehicle file is prefilled')
    fill(g, {'fault': 'تسريب زيت'}); g.click('.modal [data-save]'); g.wait_for_selector('.mr-sheet'); g.wait_for_timeout(300)
    t = g.inner_text('body')
    ok('MR-000003' in g.inner_text('.mr-sheet') and 'MR-000001' in t, 'request page lists the vehicle previous requests')
    ok(sql('select count(*) from maintenance_requests') == '3', 'previous requests kept when creating a new one')

    # ---------- list, search, filters
    g.goto(BASE + '#/requests'); g.wait_for_selector('.list-body')
    ok(g.locator('.list-body tbody tr').count() == 3, 'list shows 3 requests')
    ok('تاريخ آخر تغيير' in g.inner_text('.list-body thead') and '2026/09/01' in g.inner_text('.list-body'), 'list shows «تاريخ آخر تغيير» column')
    ok('لم يُعتمد' in g.inner_text('.list-summary'), 'list warns about requests not approved / not finished')
    g.select_option('#page select[name=extra]', 'completed'); g.wait_for_timeout(200)
    ok(g.locator('.list-body tbody tr').count() == 1, 'filter by status')
    g.select_option('#page select[name=extra]', ''); g.fill('#page input[name=q]', 'تمديد الري'); g.wait_for_timeout(200)
    ok(g.locator('.list-body tbody tr').count() == 1, 'list search by project text')
    g.fill('#page input[name=q]', ''); g.wait_for_timeout(100)
    g.screenshot(path=f'{OUT}/r-list-desktop.png', full_page=True)
    for q, what, n in [('MR-000002', 'request number', 1), ('ج 1234', 'plate', 2), ('تمديد الري', 'project text', 1), ('الغملاس', 'project (site) name', 3)]:
        g.fill('form.search input', q); g.press('form.search input', 'Enter'); g.wait_for_selector('text=نتائج البحث'); g.wait_for_timeout(400)
        rows = g.locator('section.followups:has-text("طلبات الصيانة") tbody tr').count()
        ok(rows == n, f'global search by {what} finds {n} request(s) ({rows})')

    # ---------- alerts: open requests appear in the bell and dashboard
    g.goto(BASE + '#/dashboard'); g.wait_for_selector('.followups'); g.wait_for_timeout(300)
    ok('طلب صيانة' in g.inner_text('.followups'), 'dashboard follow-ups list the open request')
    g.click('.followups a[href*="#/requests/"] >> nth=0'); g.wait_for_selector('.mr-sheet')
    ok('MR-000003' in g.inner_text('.mr-sheet'), 'alert opens the request')
    g.wait_for_timeout(21000); g.goto(BASE + '#/requests'); g.wait_for_selector('.list-body'); g.wait_for_timeout(600); g.click('[data-act=bell]'); g.wait_for_timeout(300)
    ok('طلب صيانة' in g.inner_text('.pop-bell'), 'bell shows the request alert')

    # ---------- old maintenance screen still works
    g.goto(BASE + '#/requests'); g.wait_for_selector('.list-summary a[href="#/maintenance"]'); g.click('.list-summary a[href="#/maintenance"]')
    g.wait_for_selector('h1:has-text("سجل فواتير الصيانة")'); g.wait_for_selector('.list-body'); ok('MNT-000001' in g.inner_text('.list-body') and 'MNT-000002' in g.inner_text('.list-body'), 'old maintenance records page still reachable with all records')

    # ---------- Excel export
    g.goto(BASE + '#/excel'); g.wait_for_selector('.exp-grid')
    with g.expect_download() as dl: g.click('[data-exp=requests]')
    path = f'{OUT}/requests.xlsx'; dl.value.save_as(path); ws = openpyxl.load_workbook(path).active
    ok(ws.title == 'طلبات الصيانة' and ws.max_row == 4 and ws.cell(1, 1).value == 'رقم الطلب', 'Excel export of requests (3 rows)')
    c.close()

    # ---------- mobile screens
    c, m = ctx(b, 390); login(m)
    m.goto(BASE + '#/requests'); m.wait_for_selector('.list-body'); m.screenshot(path=f'{OUT}/r-list-mobile.png', full_page=True)
    m.click('[data-list-add]'); m.wait_for_selector('.modal [name=fault]'); m.screenshot(path=f'{OUT}/r-form-mobile.png', full_page=True)
    (m.click('.modal [data-close]'), m.wait_for_timeout(150), m.locator('[data-yes]').count() and m.click('[data-yes]'))
    rid = sql("select id from maintenance_requests where request_no='MR-000001'")
    m.goto(BASE + f'#/requests/{rid}'); m.wait_for_selector('.mr-sheet')
    ok(m.evaluate('document.documentElement.scrollWidth <= window.innerWidth + 1'), 'mobile request page has no horizontal scroll')
    m.screenshot(path=f'{OUT}/r-page-mobile.png', full_page=True)
    m.click('.sidebar-toggle, [data-act=more]') if m.locator('[data-act=more]').count() else None
    c.close()

    # ---------- permissions
    sql("insert into user_permissions(user_id, permission) values ('a0000000-0000-0000-0000-000000000002','create'),('a0000000-0000-0000-0000-000000000002','edit')")
    c, v = ctx(b, 1280); login(v, 'viewer@test.local')
    pend = sql("select id from maintenance_requests where request_no='MR-000003'")
    v.goto(BASE + f'#/requests/{pend}'); v.wait_for_selector('.mr-sheet')
    ok(v.locator('[data-act=decide]').count() == 0 and v.locator('[data-act=edit]').count() == 1, 'clerk can edit but has no approval button')
    sql("insert into user_permissions(user_id, permission) values ('a0000000-0000-0000-0000-000000000002','approve_maintenance')")
    v.reload(); v.wait_for_selector('.mr-sheet')
    ok(v.locator('[data-act=decide]').count() == 1, 'with «اعتماد طلبات الصيانة» the approval button appears')
    c.close()
    b.close()

print(f'PASS {len(passed)}')
for x in passed: print('  ok  ', x)
print(f'FAIL {len(problems)}')
for x in problems: print('  FAIL', x)
